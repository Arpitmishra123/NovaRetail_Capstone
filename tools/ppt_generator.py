import json
import os

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.dml.color import RGBColor


# =========================================================
# FILE PATHS
# =========================================================

template_path = "assets/novaretail_template.pptx"

content_path = "output/slide_content.json"
icon_mapping_path = "output/icon_mapping.json"

output_path = "output/NovaRetail_QBR_Final.pptx"


# =========================================================
# LOAD JSON FILES
# =========================================================

with open(
    content_path,
    "r",
    encoding="utf-8"
) as file:
    slide_content = json.load(file)


with open(
    icon_mapping_path,
    "r",
    encoding="utf-8"
) as file:
    icon_mapping = json.load(file)


print("Slide content loaded successfully.")
print("Icon mapping loaded successfully.")


# =========================================================
# LOAD POWERPOINT TEMPLATE
# =========================================================

prs = Presentation(template_path)

print("PowerPoint template loaded successfully.")
print("Total slides:", len(prs.slides))


if len(prs.slides) != 5:
    raise ValueError(
        "Template must contain exactly 5 slides."
    )


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def get_shape_by_name(slide, shape_name):
    """
    Find a PowerPoint shape using its exact template name.
    """

    for shape in slide.shapes:

        if shape.name == shape_name:
            return shape

    raise ValueError(
        f"Shape not found: {shape_name}"
    )


def set_shape_text(slide, shape_name, text):
    """
    Replace text inside a named PowerPoint shape.
    """

    shape = get_shape_by_name(
        slide,
        shape_name
    )

    if not shape.has_text_frame:
        raise ValueError(
            f"Shape has no text frame: {shape_name}"
        )

    shape.text_frame.clear()

    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = str(text)

    shape.text_frame.word_wrap = True

    shape.text_frame.auto_size = (
        MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    )


def set_text_color(slide, shape_name, rgb_color):
    """
    Change text color inside a named PowerPoint shape.
    """

    shape = get_shape_by_name(
        slide,
        shape_name
    )

    if not shape.has_text_frame:
        raise ValueError(
            f"Shape has no text frame: {shape_name}"
        )

    for paragraph in shape.text_frame.paragraphs:

        for run in paragraph.runs:

            run.font.color.rgb = RGBColor(
                *rgb_color
            )


def remove_shape_if_exists(slide, shape_name):
    """
    Remove an instruction label or placeholder shape.
    """

    for shape in list(slide.shapes):

        if shape.name == shape_name:

            shape_element = shape._element

            shape_element.getparent().remove(
                shape_element
            )

            return True

    return False


def replace_shape_with_picture(
    slide,
    shape_name,
    image_path
):
    """
    Read the placeholder position and size,
    remove the placeholder,
    and insert an image in exactly
    the same location.
    """

    if not os.path.exists(image_path):

        raise FileNotFoundError(
            f"Image not found: {image_path}"
        )

    placeholder = get_shape_by_name(
        slide,
        shape_name
    )

    left = placeholder.left
    top = placeholder.top
    width = placeholder.width
    height = placeholder.height

    placeholder_element = placeholder._element

    placeholder_element.getparent().remove(
        placeholder_element
    )

    slide.shapes.add_picture(
        image_path,
        left,
        top,
        width,
        height
    )


def format_bullet_list(items):
    """
    Convert a list of insights/actions
    into PowerPoint-ready bullet text.
    """

    clean_items = []

    for item in items:

        item_text = str(item).strip()

        if item_text:

            clean_items.append(
                f"• {item_text}"
            )

    return "\n".join(
        clean_items
    )


def get_icon_path(mapping_key):
    """
    Get full icon path from icon_mapping.json.
    """

    icon_filename = icon_mapping.get(
        mapping_key
    )

    if not icon_filename:

        raise ValueError(
            f"Icon mapping missing: {mapping_key}"
        )

    return os.path.join(
        "icons",
        icon_filename
    )


# =========================================================
# SLIDE 1: EXECUTIVE SUMMARY
# =========================================================

slide_1 = prs.slides[0]
content_1 = slide_content["slide_1"]


set_shape_text(
    slide_1,
    "deck_title",
    content_1.get(
        "title",
        "NovaRetail Quarterly Business Review"
    )
)

set_shape_text(
    slide_1,
    "subtitle_1",
    content_1.get(
        "review_period",
        ""
    )
)


set_shape_text(
    slide_1,
    "kpi1_value",
    content_1["kpi_1"]["value"]
)

set_shape_text(
    slide_1,
    "kpi1_label",
    content_1["kpi_1"]["label"]
)


set_shape_text(
    slide_1,
    "kpi2_value",
    content_1["kpi_2"]["value"]
)

set_shape_text(
    slide_1,
    "kpi2_label",
    content_1["kpi_2"]["label"]
)


set_shape_text(
    slide_1,
    "kpi3_value",
    content_1["kpi_3"]["value"]
)

set_shape_text(
    slide_1,
    "kpi3_label",
    content_1["kpi_3"]["label"]
)


set_shape_text(
    slide_1,
    "summary_text",
    content_1["executive_summary"]
)


replace_shape_with_picture(
    slide_1,
    "icon_1",
    get_icon_path(
        "slide_1_executive_summary"
    )
)

remove_shape_if_exists(
    slide_1,
    "icon_1_label"
)

# Improve text contrast on dark background

WHITE = (255, 255, 255)
LIGHT_GRAY = (220, 225, 235)

set_text_color(
    slide_1,
    "deck_title",
    WHITE
)

set_text_color(
    slide_1,
    "subtitle_1",
    LIGHT_GRAY
)

set_text_color(
    slide_1,
    "kpi1_value",
    WHITE
)

set_text_color(
    slide_1,
    "kpi1_label",
    LIGHT_GRAY
)

set_text_color(
    slide_1,
    "kpi2_value",
    WHITE
)

set_text_color(
    slide_1,
    "kpi2_label",
    LIGHT_GRAY
)

set_text_color(
    slide_1,
    "kpi3_value",
    WHITE
)

set_text_color(
    slide_1,
    "kpi3_label",
    LIGHT_GRAY
)

set_text_color(
    slide_1,
    "summary_label",
    WHITE
)

set_text_color(
    slide_1,
    "summary_text",
    WHITE
)

print("Slide 1 completed.")


# =========================================================
# SLIDE 2: REGIONAL PERFORMANCE
# =========================================================

slide_2 = prs.slides[1]
content_2 = slide_content["slide_2"]


set_shape_text(
    slide_2,
    "title_2",
    content_2["headline"]
)


set_shape_text(
    slide_2,
    "insights_performance",
    format_bullet_list(
        content_2["insights"]
    )
)


replace_shape_with_picture(
    slide_2,
    "chart_performance",
    "charts/slide_2_revenue.png"
)

remove_shape_if_exists(
    slide_2,
    "chart_performance_label"
)


replace_shape_with_picture(
    slide_2,
    "icon_2",
    get_icon_path(
        "slide_2_regional_performance"
    )
)

remove_shape_if_exists(
    slide_2,
    "icon_2_label"
)


print("Slide 2 completed.")


# =========================================================
# SLIDE 3: CUSTOMER EXPERIENCE
# =========================================================

slide_3 = prs.slides[2]
content_3 = slide_content["slide_3"]


set_shape_text(
    slide_3,
    "title_3",
    content_3["headline"]
)


set_shape_text(
    slide_3,
    "insights_customer",
    format_bullet_list(
        content_3["insights"]
    )
)


replace_shape_with_picture(
    slide_3,
    "chart_customer",
    "charts/slide_3_nps.png"
)

remove_shape_if_exists(
    slide_3,
    "chart_customer_label"
)


replace_shape_with_picture(
    slide_3,
    "icon_3",
    get_icon_path(
        "slide_3_customer_experience"
    )
)

remove_shape_if_exists(
    slide_3,
    "icon_3_label"
)


print("Slide 3 completed.")


# =========================================================
# SLIDE 4: OPERATIONAL RISK
# =========================================================

slide_4 = prs.slides[3]
content_4 = slide_content["slide_4"]


set_shape_text(
    slide_4,
    "title_4",
    content_4["headline"]
)


set_shape_text(
    slide_4,
    "insights_risk",
    format_bullet_list(
        content_4["insights"]
    )
)


replace_shape_with_picture(
    slide_4,
    "chart_risk",
    "charts/slide_4_stockout.png"
)

remove_shape_if_exists(
    slide_4,
    "chart_risk_label"
)


replace_shape_with_picture(
    slide_4,
    "icon_4",
    get_icon_path(
        "slide_4_operational_risk"
    )
)

remove_shape_if_exists(
    slide_4,
    "icon_4_label"
)


print("Slide 4 completed.")


# =========================================================
# SLIDE 5: RECOMMENDED ACTIONS
# =========================================================

slide_5 = prs.slides[4]
content_5 = slide_content["slide_5"]


set_shape_text(
    slide_5,
    "title_5",
    content_5.get(
        "headline",
        "Recommended Actions & Next Steps"
    )
)


set_shape_text(
    slide_5,
    "recommendations_text",
    format_bullet_list(
        content_5["actions"]
    )
)


set_shape_text(
    slide_5,
    "closing_statement",
    content_5[
        "forward_looking_statement"
    ]
)


replace_shape_with_picture(
    slide_5,
    "icon_5",
    get_icon_path(
        "slide_5_recommended_actions"
    )
)

remove_shape_if_exists(
    slide_5,
    "icon_5_label"
)


# Improve text contrast on dark background

set_text_color(
    slide_5,
    "title_5",
    WHITE
)

set_text_color(
    slide_5,
    "recommendations_text",
    WHITE
)

set_text_color(
    slide_5,
    "closing_statement",
    WHITE
)

print("Slide 5 completed.")


# =========================================================
# SAVE FINAL PRESENTATION
# =========================================================

prs.save(output_path)


print("\nSUCCESS!")
print("Final PowerPoint created:")
print(output_path)