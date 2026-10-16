from pptx import Presentation

template_path = "assets/novaretail_template.pptx"

prs = Presentation(template_path)

print("PowerPoint template opened successfully.")
print("Total slides:", len(prs.slides))

print("\nPresentation Size:")
print("Width:", prs.slide_width)
print("Height:", prs.slide_height)

print("\n===== SLIDE SHAPES =====")

for slide_number, slide in enumerate(prs.slides, start=1):

    print(f"\n--- SLIDE {slide_number} ---")

    for shape in slide.shapes:

        print(
            "Name:", shape.name,
            "| Type:", shape.shape_type,
            "| Left:", shape.left,
            "| Top:", shape.top,
            "| Width:", shape.width,
            "| Height:", shape.height
        )

        if shape.has_text_frame:
            text = shape.text.strip()

            if text:
                print("   Text:", repr(text))


print("\n===== SIMPLE SHAPE LIST =====")

for slide_number, slide in enumerate(prs.slides, start=1):

    print(f"\nSLIDE {slide_number}")

    for shape in slide.shapes:

        print("Shape:", shape.name)

        if shape.has_text_frame:

            text = shape.text.strip()

            if text:
                print("Text:", repr(text))


